"""Turn uploaded files (PDF, DOCX, TXT, Markdown) into raw text. Fully local."""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import PurePath

from ..config import MAX_UPLOAD_MB
from ..errors import ExtractionError


@dataclass
class ExtractedDocument:
    name: str
    kind: str  # pdf | docx | txt | md
    text: str
    pages: int | None = None
    warnings: list[str] = field(default_factory=list)

    @property
    def char_count(self) -> int:
        return len(self.text)


def detect_kind(filename: str) -> str:
    suffix = PurePath(filename).suffix.lower().lstrip(".")
    return {"markdown": "md", "text": "txt"}.get(suffix, suffix)


def extract(filename: str, data: bytes) -> ExtractedDocument:
    """Extract text from ``data``. Raises :class:`ExtractionError` with a friendly message."""
    if len(data) > MAX_UPLOAD_MB * 1024 * 1024:
        raise ExtractionError(f"This file is larger than {MAX_UPLOAD_MB} MB.")
    kind = detect_kind(filename)
    readers = {"pdf": _read_pdf, "docx": _read_docx, "txt": _read_text, "md": _read_text}
    reader = readers.get(kind)
    if reader is None:
        raise ExtractionError("Unsupported file type. Use PDF, DOCX, TXT or Markdown.")
    doc = reader(filename, data, kind)
    if not doc.text.strip():
        raise ExtractionError("No readable text was found in this file.")
    return doc


def _decode(data: bytes) -> tuple[str, list[str]]:
    warnings: list[str] = []
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16"), warnings
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding), warnings
        except UnicodeDecodeError:
            continue
    warnings.append("The text encoding was unclear; some characters may look wrong.")
    return data.decode("latin-1", errors="replace"), warnings


def _read_text(filename: str, data: bytes, kind: str) -> ExtractedDocument:
    text, warnings = _decode(data)
    return ExtractedDocument(filename, kind, text, warnings=warnings)


def _read_pdf(filename: str, data: bytes, kind: str) -> ExtractedDocument:
    try:
        from pypdf import PdfReader  # type: ignore
        from pypdf.errors import PdfReadError  # type: ignore
    except ImportError as exc:
        raise ExtractionError("PDF support needs 'pypdf'. Run: uv pip install pypdf") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise ExtractionError("This PDF is password-protected.")
            except ExtractionError:
                raise
            except Exception as exc:
                raise ExtractionError("This PDF is password-protected.") from exc
        pages = [(page.extract_text() or "") for page in reader.pages]
    except ExtractionError:
        raise
    except PdfReadError as exc:
        raise ExtractionError("This PDF appears to be damaged and could not be read.") from exc
    except Exception as exc:
        raise ExtractionError(f"Could not read this PDF: {exc}") from exc

    warnings: list[str] = []
    empty = sum(1 for p in pages if not p.strip())
    if pages and empty == len(pages):
        raise ExtractionError(
            "This PDF has no selectable text (it looks scanned). OCR is not included in this version."
        )
    if empty:
        warnings.append(f"{empty} of {len(pages)} pages had no extractable text (images or scans).")
    return ExtractedDocument(filename, kind, "\n\n".join(pages), pages=len(pages), warnings=warnings)


def _read_docx(filename: str, data: bytes, kind: str) -> ExtractedDocument:
    try:
        from docx import Document  # type: ignore
        from docx.table import Table  # type: ignore
        from docx.text.paragraph import Paragraph  # type: ignore
    except ImportError as exc:
        raise ExtractionError("DOCX support needs 'python-docx'. Run: uv pip install python-docx") from exc
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("This DOCX file could not be opened. Is it a valid Word document?") from exc

    lines: list[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                lines.append("")
                continue
            style = (paragraph.style.name or "") if paragraph.style is not None else ""
            if (style.startswith("Heading") or style == "Title") and text[-1] not in ".!?:":
                text += "."
            lines.append(text)
        elif tag == "tbl":
            table = Table(child, document)
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(", ".join(cells) + ".")
            lines.append("")
    return ExtractedDocument(filename, kind, "\n".join(lines))
